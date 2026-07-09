"""利用者(患者)の Excel 一括取込.

1 行 = 1 利用者。列: 利用者コード / 氏名 / 性別 / 生年月日 / ペースメーカー / 服薬情報 / 備考。
列名はゆらぎを吸収する。空欄は「未入力」として扱い、既存利用者の更新時は上書きしない。
"""

from __future__ import annotations

import io
import re
import sqlite3
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Optional

import pandas as pd

from .. import repo
from ..domain import parse_date

CODE_HEADERS = ["利用者コード", "利用者id", "利用者番号", "コード", "id", "利用者"]
NAME_HEADERS = ["氏名", "名前", "利用者名", "name"]
SEX_HEADERS = ["性別", "sex"]
BIRTH_HEADERS = ["生年月日", "誕生日", "生年月", "birth"]
PM_HEADERS = ["ペースメーカー", "ペースメーカ", "pacemaker", "pm"]
MED_HEADERS = ["服薬情報", "服薬", "薬剤", "投薬"]
NOTE_HEADERS = ["備考", "メモ", "note"]

TEMPLATE_HEADERS = ["利用者コード", "氏名", "性別", "生年月日", "ペースメーカー", "服薬情報", "備考"]

_SEX_MAP = {
    "男": "M", "男性": "M", "m": "M", "male": "M", "1": "M",
    "女": "F", "女性": "F", "f": "F", "female": "F", "2": "F",
}
_PM_TRUE = {"1", "有", "あり", "有り", "yes", "y", "true", "○", "◯", "o"}
_PM_FALSE = {"0", "無", "なし", "無し", "no", "n", "false", "×", "x"}

# Excel シリアル値の起点(1900年システムの既知のズレを含む)
_EXCEL_EPOCH = datetime(1899, 12, 30)


def _norm(s) -> str:
    if s is None:
        return ""
    return re.sub(r"\s+", "", str(s)).strip().lower()


def _is_blank(v) -> bool:
    """None / NaN / NaT / 空文字 を空とみなす."""
    if v is None:
        return True
    try:
        if pd.isna(v):  # NaN, NaT。pd.NaT は datetime のサブクラスなので先に弾く
            return True
    except (TypeError, ValueError):
        pass  # 配列等、isna がスカラーを返さないケース
    return str(v).strip() == ""


def _cell(v) -> str:
    """セル値 → 前後空白を除いた文字列。空/NaN/NaT は ''。"""
    if _is_blank(v):
        return ""
    return str(v).strip()


def _parse_sex(v) -> Optional[str]:
    t = _norm(v)
    if t == "":
        return None
    # "1.0" のような数値セル
    t = re.sub(r"\.0$", "", t)
    return _SEX_MAP.get(t)


def _parse_pacemaker(v) -> Optional[int]:
    t = _norm(v)
    t = re.sub(r"\.0$", "", t)
    if t == "":
        return None
    if t in _PM_TRUE:
        return 1
    if t in _PM_FALSE:
        return 0
    return None


def _parse_birth(v) -> tuple[Optional[str], bool]:
    """(ISO文字列 or None, 解釈できたか) を返す。未入力は (None, True)。"""
    if _is_blank(v):
        return None, True
    if isinstance(v, (datetime, date)):
        d = parse_date(v)
        return (d.isoformat() if d else None), d is not None
    raw = _cell(v)
    # Excel シリアル値(日付書式でないセル)
    if re.fullmatch(r"\d{5}(\.0)?", raw):
        try:
            d = (_EXCEL_EPOCH + timedelta(days=int(float(raw)))).date()
            return d.isoformat(), True
        except (ValueError, OverflowError):
            return None, False
    d = parse_date(raw)
    return (d.isoformat() if d else None), d is not None


@dataclass
class PatientRow:
    row_no: int
    code: str
    name: str = ""
    sex: Optional[str] = None
    birth_date: Optional[str] = None
    pacemaker: Optional[int] = None
    med_info: str = ""
    note: str = ""
    status: str = "new"          # new / update / error
    errors: list[str] = field(default_factory=list)
    patient_id: Optional[int] = None


@dataclass
class PatientImportPreview:
    rows: list[PatientRow] = field(default_factory=list)
    missing_headers: list[str] = field(default_factory=list)

    @property
    def new_rows(self) -> list[PatientRow]:
        return [r for r in self.rows if r.status == "new"]

    @property
    def update_rows(self) -> list[PatientRow]:
        return [r for r in self.rows if r.status == "update"]

    @property
    def error_rows(self) -> list[PatientRow]:
        return [r for r in self.rows if r.status == "error"]

    @property
    def has_errors(self) -> bool:
        return bool(self.error_rows)


def read_workbook(path: str | Path) -> pd.DataFrame:
    p = Path(path)
    engine = "xlrd" if p.suffix.lower() == ".xls" else "openpyxl"
    df = pd.read_excel(p, engine=engine, header=0, dtype=object)
    return df.dropna(axis=1, how="all").dropna(axis=0, how="all")


def _find_col(headers, candidates) -> Optional[str]:
    norm = [(h, _norm(h)) for h in headers]
    for cand in candidates:
        for h, nh in norm:
            if nh == cand:
                return h
    for cand in candidates:
        for h, nh in norm:
            if cand in nh:
                return h
    return None


def analyze(conn: sqlite3.Connection, path: str | Path) -> PatientImportPreview:
    df = read_workbook(path)
    headers = list(df.columns)

    c_code = _find_col(headers, CODE_HEADERS)
    c_name = _find_col(headers, NAME_HEADERS)
    c_sex = _find_col(headers, SEX_HEADERS)
    c_birth = _find_col(headers, BIRTH_HEADERS)
    c_pm = _find_col(headers, PM_HEADERS)
    c_med = _find_col(headers, MED_HEADERS)
    c_note = _find_col(headers, NOTE_HEADERS)

    preview = PatientImportPreview()
    if c_code is None:
        preview.missing_headers.append("利用者コード")
        return preview

    seen: dict[str, int] = {}
    for i, (_, r) in enumerate(df.iterrows(), start=2):  # 2 = ヘッダ次の行番号
        code = _cell(r.get(c_code))
        if re.fullmatch(r"\d+\.0", code):
            code = code[:-2]
        if code == "":
            continue  # 空行はスキップ

        pr = PatientRow(row_no=i, code=code)
        pr.name = _cell(r.get(c_name)) if c_name else ""
        pr.med_info = _cell(r.get(c_med)) if c_med else ""
        pr.note = _cell(r.get(c_note)) if c_note else ""

        if c_sex:
            raw_sex = _cell(r.get(c_sex))
            pr.sex = _parse_sex(r.get(c_sex))
            if raw_sex and pr.sex is None:
                pr.errors.append(f"性別を解釈できません: {raw_sex}")

        if c_birth:
            pr.birth_date, ok = _parse_birth(r.get(c_birth))
            if not ok:
                pr.errors.append(f"生年月日を解釈できません: {_cell(r.get(c_birth))}")

        if c_pm:
            raw_pm = _cell(r.get(c_pm))
            pr.pacemaker = _parse_pacemaker(r.get(c_pm))
            if raw_pm and pr.pacemaker is None:
                pr.errors.append(f"ペースメーカーを解釈できません: {raw_pm}")

        if code in seen:
            pr.errors.append(f"ファイル内でコードが重複しています(行 {seen[code]})")
        else:
            seen[code] = i

        existing = repo.get_patient_by_code(conn, code)
        if existing:
            pr.patient_id = existing["id"]

        if pr.errors:
            pr.status = "error"
        else:
            pr.status = "update" if existing else "new"

        preview.rows.append(pr)

    return preview


def commit(
    conn: sqlite3.Connection,
    path: str | Path,
    update_existing: bool = True,
) -> dict:
    """取込を確定。既存利用者は update_existing=True のとき、
    ファイルに値がある項目だけを上書きする(空欄は既存値を保持)。"""
    preview = analyze(conn, path)
    created = updated = skipped = 0

    for pr in preview.rows:
        if pr.status == "error":
            skipped += 1
            continue
        if pr.status == "new":
            repo.create_patient(
                conn, code=pr.code, name=(pr.name or pr.code), sex=pr.sex,
                birth_date=pr.birth_date, pacemaker=(pr.pacemaker or 0),
                note=pr.note, med_info=pr.med_info,
            )
            created += 1
        else:
            if not update_existing:
                skipped += 1
                continue
            cur = repo.get_patient(conn, pr.patient_id)
            repo.update_patient(
                conn, pr.patient_id,
                name=pr.name or cur["name"],
                sex=pr.sex if pr.sex is not None else cur["sex"],
                birth_date=pr.birth_date or cur["birth_date"],
                pacemaker=pr.pacemaker if pr.pacemaker is not None else cur["pacemaker"],
                note=pr.note or cur["note"],
                med_info=pr.med_info or cur["med_info"],
            )
            updated += 1

    return {"created": created, "updated": updated, "skipped": skipped}


def build_template() -> bytes:
    """入力用テンプレート xlsx を生成する."""
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = "利用者"
    ws.append(TEMPLATE_HEADERS)
    for c in ws[1]:
        c.font = Font(bold=True)
    ws.append(["U001", "山田太郎", "男", "1950-05-01", "無", "", ""])
    ws.append(["U002", "佐藤花子", "女", "1948-11-23", "有", "血糖降下薬", "要見守り"])
    widths = [14, 16, 8, 14, 16, 20, 20]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
