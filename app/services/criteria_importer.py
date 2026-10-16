"""評価基準 Excel(「評価項目 改訂案」形式)の取込.

Excel の構成:
  - 1 シート = 1 つの(年代, ペースメーカー有無)。シート名は「50代以下」「60代」
    「60代　ペースメーカー」のような形式(空白はゆらぐ)。それ以外のシートは対象外。
  - 見出し行「評価項目 | 評価基準内容 | 相対評価 | 担当者 | 評価コメント | 評価項目説明 | 備考」
    の下に、項目名の行から空行までが 1 項目のブロックとして縦に並ぶ。
  - 評価基準内容 = 区間下限値、相対評価 = 1〜5。定性項目(FIM(社会性)等)は相対評価欄が
    文言なので、評価基準内容のコード値(1〜5)をそのまま評価とする。
  - 「男　BMI」のような接頭辞(または項目名行の評価基準内容欄の 男/女)で性別条件になる。

取込はシートごとに読んだ基準を項目ごとに集め、``normalize`` で冗長な展開を畳んでから
(全シート同一なら条件なし、PM 有無で同一なら pacemaker=NULL)、対象項目の基準を置換する。
Excel 側の判定(VLOOKUP 近似一致 = 測定値以下で最大の下限値)は評価エンジンと同じ。
"""

from __future__ import annotations

import re
import sqlite3
import unicodedata
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Optional

from .. import config, repo
from ..domain import AGE_BAND_LABELS, age_to_band

# Excel の項目名(_key で正規化) -> アプリの評価項目名
EXCEL_ITEM_MAP = {
    "hba1c": "HbA1c",
    "血圧(mmhg)": "収縮期血圧",
    "bmi": "BMI",
    "smi": "SMI",
    "喫食率(割)": "喫食率",
    "体重減少率(%)": "体重増減率",
    "骨密度(腰)": "骨密度(腰椎)",
    "骨密度(大腿骨)": "骨密度(大腿骨)",
    "tug(秒)": "TUG",
    "握力(kg)": "握力",
    "frt(cm)": "FRT",
    "バーセル": "Barthel Index",
    "mini-cog": "mini-Cog",
    "fim(社会性)": "FIM",
    "摂食嚥下機能": "嚥下機能",
    "口腔内": "口腔機能",
    "聴力(1000hz)": "聴力1000Hz",
    "聴力(4000hz)": "聴力4000Hz",
    "視力": "視力",
    "下腿周囲長(cm)": "下腿周囲長",
}

# アプリに既定で無い項目。DB に無ければ「無効」で作成する(有効化は基準マスタ画面で)。
# after: 表示順をこの項目に揃える
NEW_ITEMS = {
    "視力": dict(unit="", category="感覚機能", direction="higher_better", in_radar=0,
                 source_column="視力", after="聴力4000Hz"),
    "下腿周囲長": dict(unit="cm", category="身体機能", direction="higher_better", in_radar=0,
                    source_column="下腿周囲長", after="SMI"),
}

# このブロックは評価基準ではない(総合評価文言は帳票設定で管理する)
NON_CRITERIA_NAMES = {"総合評価"}

SEX_PREFIX = {"男": "M", "女": "F"}
SEX_LABEL = {"M": "男", "F": "女"}

_SHEET_RE = re.compile(r"^(\d+)代(以下|以上)?(ペースメーカー|PM)?$", re.IGNORECASE)
# 見出しの補助ラベル(「６０代」「(高音域)」など)。項目ではないので読み飛ばす。
_LABEL_RE = re.compile(r"^(\d+代|\(.*\))$")

# (threshold, score, comment, description)
Row = tuple[float, int, Optional[str], Optional[str]]


def _nfkc(s) -> str:
    return unicodedata.normalize("NFKC", str(s)) if s is not None else ""


def _clean(s) -> str:
    """NFKC + 空白除去(全角空白・改行を含む)."""
    return re.sub(r"\s+", "", _nfkc(s))


def _key(name: str) -> str:
    return _clean(name).lower()


def _text(v) -> Optional[str]:
    if v is None:
        return None
    s = str(v).strip()
    return s or None


def _number(v) -> Optional[float]:
    if isinstance(v, bool) or v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    try:
        return float(_clean(v))
    except ValueError:
        return None


def sheet_condition(sheet_name: str) -> Optional[tuple[str, bool]]:
    """シート名 → (年代バンド, ペースメーカー有のシートか)。対象外シートは None."""
    m = _SHEET_RE.match(_clean(sheet_name))
    if not m:
        return None
    return age_to_band(int(m.group(1))), bool(m.group(3))


def split_sex(name: str) -> tuple[Optional[str], str]:
    n = _clean(name)
    for pref, code in SEX_PREFIX.items():
        if n.startswith(pref) and len(n) > 1:
            return code, n[1:]
    return None, n


# --- シート解析 --------------------------------------------------------------

_HEADERS = {
    "item": ["評価項目"],
    "threshold": ["評価基準", "基準値", "閾値"],
    "score": ["相対評価"],
    "comment": ["評価コメント", "コメント"],
    "desc": ["評価項目説明", "説明"],
}
# 見出しが見つからないときの既定列(A,B,C,E,F)
_DEFAULT_COLS = {"item": 0, "threshold": 1, "score": 2, "comment": 4, "desc": 5}


def _find_header(rows: list[tuple]) -> tuple[int, dict[str, int]]:
    """見出し行の index と列位置を返す。見つからなければ (-1, 既定列)."""
    for i, row in enumerate(rows[:15]):
        cells = [_clean(v) for v in row]
        if "評価項目" in cells and any("相対評価" in c for c in cells):
            cols = dict(_DEFAULT_COLS)
            for k, cands in _HEADERS.items():
                for j, c in enumerate(cells):
                    if c and any(cand in c for cand in cands):
                        cols[k] = j
                        break
            return i, cols
    return -1, dict(_DEFAULT_COLS)


def _cell(row: tuple, idx: int):
    return row[idx] if idx < len(row) else None


def parse_sheet(rows: list[tuple], sheet_label: str = "") -> tuple[dict, list[str]]:
    """シートの行(値のタプル列) → {(項目名, sex): [Row...]} と警告."""
    warnings: list[str] = []
    header_idx, cols = _find_header(rows)
    if header_idx < 0:
        return {}, [f"{sheet_label}: 見出し行(評価項目/相対評価)が見つからないため読み飛ばしました"]

    result: dict[tuple[str, Optional[str]], list[Row]] = {}
    current: Optional[list[Row]] = None   # 取込中ブロックの行リスト
    current_name = ""
    skipping = False                      # 総合評価など基準でないブロック

    for row in rows[header_idx + 1:]:
        a = _text(_cell(row, cols["item"]))
        thr_raw = _cell(row, cols["threshold"])
        score_raw = _cell(row, cols["score"])
        comment = _text(_cell(row, cols["comment"]))
        desc = _text(_cell(row, cols["desc"]))

        if a is None and _text(thr_raw) is None and _text(score_raw) is None and comment is None:
            current, skipping = None, False   # 空行でブロック終了
            continue

        if a is not None:
            sex, base = split_sex(a)
            if base in NON_CRITERIA_NAMES:
                current, skipping = None, True
                continue
            known = _key(base) in EXCEL_ITEM_MAP
            if current is None and _LABEL_RE.match(base):
                continue  # ブロック前の補助ラベル
            if current is not None and not known:
                # 取込中ブロック内の未知の文字列は「4000Hｚ」のような注記として同じ項目の行に含める。
                # 空行を挟まずに書かれた未知の項目を誤って混ぜていないか確認できるよう警告する。
                warnings.append(f"{sheet_label}: 「{a}」の行は直前の「{current_name}」の一部として読みました")
            if current is None or known:
                # 新しい項目ブロック
                thr_text = _clean(thr_raw)
                if sex is None and thr_text in SEX_PREFIX:
                    sex = SEX_PREFIX[thr_text]   # 項目名行の評価基準内容欄に性別がある書式
                key = (base, sex)
                if key in result:
                    warnings.append(f"{sheet_label}: 「{a}」が重複しているため2つ目以降を無視しました")
                    current, skipping = None, True
                    continue
                current, skipping = [], False
                current_name = a
                result[key] = current

        if current is None:
            continue

        thr = _number(thr_raw)
        score_num = _number(score_raw)
        if thr is None:
            if score_num is not None or comment:
                warnings.append(
                    f"{sheet_label}: 「{current_name}」の評価基準内容が空欄/数値でない行を除外しました"
                    f"(相対評価 {_text(score_raw) or '空欄'})"
                )
            continue  # 「上」などの小見出し
        if score_num is not None:
            if not score_num.is_integer():
                warnings.append(f"{sheet_label}: 「{current_name}」の相対評価 {score_num:g} は整数でないため除外しました")
                continue
            score = int(score_num)
        elif _text(score_raw) is not None and thr.is_integer():
            score = int(thr)  # 定性項目: 相対評価欄は文言、コード値が評価
        else:
            warnings.append(f"{sheet_label}: 「{current_name}」の相対評価が読めない行を除外しました(基準値 {thr:g})")
            continue
        if not 1 <= score <= 5:
            warnings.append(f"{sheet_label}: 「{current_name}」の相対評価 {score} は 1〜5 の範囲外のため除外しました")
            continue
        if any(r[0] == thr for r in current):
            warnings.append(f"{sheet_label}: 「{current_name}」の評価基準内容 {thr:g} が重複しているため2つ目以降を除外しました")
            continue
        current.append((thr, score, comment, desc))

    return {k: v for k, v in result.items() if v}, warnings


def _read_workbook(path: Path) -> dict[str, list[tuple]]:
    from openpyxl import load_workbook

    if Path(path).suffix.lower() != ".xlsx":
        raise ValueError("評価基準の取込は .xlsx 形式のみ対応しています。Excel で .xlsx 形式に保存し直してください。")
    wb = load_workbook(path, data_only=True, read_only=True)
    try:
        return {ws.title: [tuple(r) for r in ws.iter_rows(values_only=True)] for ws in wb.worksheets}
    finally:
        wb.close()


@dataclass
class ParsedWorkbook:
    # (項目名, sex) -> {(age_band, pacemaker): [Row...]}
    data: dict[tuple[str, Optional[str]], dict[tuple, list[Row]]]
    all_conds: set[tuple]
    display_names: dict[str, str]          # 正規化名 -> Excel 上の表記
    skipped_sheets: list[str]
    warnings: list[str]


def collect(sheets: dict[str, list[tuple]]) -> ParsedWorkbook:
    targets: dict[tuple, str] = {}
    skipped: list[str] = []
    warnings: list[str] = []
    by_band: dict[str, list[bool]] = defaultdict(list)
    parsed_conds: dict[str, tuple[str, bool]] = {}
    for name in sheets:
        cond = sheet_condition(name)
        if cond is None:
            skipped.append(name)
            continue
        parsed_conds[name] = cond
        by_band[cond[0]].append(cond[1])

    for name, (band, is_pm) in parsed_conds.items():
        # PM 有のシートがある年代は PM 無/有 を区別し、無い年代(50代以下)は PM 条件なし
        has_pm_variant = True in by_band[band]
        pm = (1 if is_pm else 0) if has_pm_variant else None
        if (band, pm) in targets:
            warnings.append(f"シート「{name}」は「{targets[(band, pm)]}」と同じ条件のため読み飛ばしました")
            skipped.append(name)
            continue
        targets[(band, pm)] = name

    data: dict = defaultdict(dict)
    display: dict[str, str] = {}
    for cond, sheet_name in sorted(targets.items(), key=lambda kv: (kv[0][0], kv[0][1] is not None, kv[0][1] or 0)):
        parsed, w = parse_sheet(sheets[sheet_name], sheet_label=sheet_name.strip())
        warnings.extend(w)
        for (base, sex), rows in parsed.items():
            data[(base, sex)][cond] = sorted(rows, key=lambda r: r[0])
            display.setdefault(_key(base), base)
    return ParsedWorkbook(dict(data), set(targets), display, skipped, warnings)


# --- 冗長展開の解消 -----------------------------------------------------------

def _same(a: list, b: list) -> bool:
    """閾値・評価・コメントが同じなら同一(説明文の差は無視)."""
    return [r[:3] for r in a] == [r[:3] for r in b]


def normalize(cond_rows: dict[tuple, list], all_conds: Optional[set] = None):
    """{(age_band, pm): rows} → [(age_band, pacemaker, rows)]。

    - 全シートに存在し全て同一 → (None, None) 1 セットに集約
    - 同一年代で PM 無/有が同一 → pacemaker=None に集約
    ``all_conds`` はブック内の全シート条件。項目が一部のシートにしか無い場合
    (SMI は PM 有シートに無い等)に、無いシートへ基準を広げないために使う。
    """
    if not cond_rows:
        return []
    all_conds = set(cond_rows) if all_conds is None else set(all_conds)
    rowsets = list(cond_rows.values())
    if set(cond_rows) == all_conds and all(_same(rowsets[0], s) for s in rowsets):
        return [(None, None, rowsets[0])]

    by_age: dict[Optional[str], dict] = defaultdict(dict)
    for (age_band, pm), rows in cond_rows.items():
        by_age[age_band][pm] = rows

    out = []
    for age_band in sorted(by_age, key=lambda a: a or ""):
        pm_map = by_age[age_band]
        sets = list(pm_map.values())
        if len(sets) > 1 and all(_same(sets[0], s) for s in sets):
            out.append((age_band, None, sets[0]))
        else:
            for pm in sorted(pm_map, key=lambda p: -1 if p is None else p):
                out.append((age_band, pm, pm_map[pm]))
    return out


# --- 取込計画(プレビュー)と適用 ------------------------------------------------

@dataclass
class CondSet:
    sex: Optional[str]
    age_band: Optional[str]
    pacemaker: Optional[int]
    rows: list[Row]

    @property
    def label(self) -> str:
        return cond_label(self.sex, self.age_band, self.pacemaker)


@dataclass
class ItemPlan:
    excel_name: str
    item_name: str
    item_id: Optional[int]          # None = 新規作成
    new_item: Optional[dict]        # 新規作成時の属性
    cond_sets: list[CondSet] = field(default_factory=list)
    current_rows: int = 0

    @property
    def is_new(self) -> bool:
        return self.item_id is None

    @property
    def row_count(self) -> int:
        return sum(len(c.rows) for c in self.cond_sets)


@dataclass
class CriteriaPreview:
    plans: list[ItemPlan]
    unknown: list[str]
    skipped_sheets: list[str]
    warnings: list[str]


def cond_label(sex, age_band, pacemaker) -> str:
    parts = [SEX_LABEL.get(sex, "")] if sex else []
    parts.append(AGE_BAND_LABELS.get(age_band, "") if age_band else "全年代")
    if pacemaker is not None:
        parts.append("PM有" if pacemaker else "PM無")
    return " ".join(p for p in parts if p)


def _resolve_item(conn: sqlite3.Connection, excel_base: str) -> tuple[Optional[sqlite3.Row], Optional[str]]:
    """Excel 項目名 → (既存項目, アプリ上の項目名)。対応先が無ければ (None, None)."""
    canonical = EXCEL_ITEM_MAP.get(_key(excel_base))
    candidates = [c for c in (canonical, excel_base, re.sub(r"\(.*\)$", "", excel_base)) if c]
    aliases = repo.alias_map(conn)
    by_key = {_key(it["name"]): it for it in repo.list_items(conn, include_inactive=True)}
    for cand in candidates:
        it = by_key.get(_key(cand))
        if it is not None:
            return it, it["name"]
        if cand in aliases:
            it = repo.get_item(conn, aliases[cand])
            return it, it["name"]
    if canonical in NEW_ITEMS:
        return None, canonical
    return None, None


def analyze(conn: sqlite3.Connection, path: Path) -> CriteriaPreview:
    parsed = collect(_read_workbook(path))

    warnings = list(parsed.warnings)
    # Excel の表記ゆれ(シートにより「TUG(秒)」と「TUG」等)は同じアプリ項目にまとめる。
    # 項目ごとに 1 計画にしないと、適用時に後の計画が先の計画の基準を消してしまう。
    resolved: dict[str, tuple[Optional[sqlite3.Row], list[str], dict]] = {}
    unknown: list[str] = []
    for (base, sex), cond_rows in parsed.data.items():
        item, item_name = _resolve_item(conn, base)
        if item_name is None:
            if base not in unknown:
                unknown.append(base)
            continue
        _item, spellings, by_sex = resolved.setdefault(item_name, (item, [], defaultdict(dict)))
        if base not in spellings:
            spellings.append(base)
        merged = by_sex[sex]
        for cond, rows in cond_rows.items():
            if cond in merged:
                warnings.append(
                    f"「{item_name}」に対応する項目が同じシートに複数あります"
                    f"({' / '.join(spellings)}、{cond_label(sex, *cond)})。先に読んだ方を使います"
                )
                continue
            merged[cond] = rows

    plans: list[ItemPlan] = []
    for item_name, (item, spellings, by_sex) in resolved.items():
        plan = ItemPlan(
            excel_name=" / ".join(spellings), item_name=item_name,
            item_id=item["id"] if item is not None else None,
            new_item=None if item is not None else NEW_ITEMS[item_name],
            current_rows=len(repo.list_criteria_for_item(conn, item["id"])) if item is not None else 0,
        )
        for sex in sorted(by_sex, key=lambda s: s or ""):
            for age_band, pm, rows in normalize(by_sex[sex], parsed.all_conds):
                plan.cond_sets.append(CondSet(sex, age_band, pm, rows))
        plans.append(plan)

    order = {it["name"]: (it["sort_order"], it["id"]) for it in repo.list_items(conn, include_inactive=True)}
    plans.sort(key=lambda p: order.get(p.item_name, order.get((p.new_item or {}).get("after"), (10**6, 0))))
    return CriteriaPreview(plans, unknown, parsed.skipped_sheets, warnings)


def backup_db(conn: sqlite3.Connection) -> Path:
    """置換前の DB を data/backups にコピーする."""
    d = config.data_dir() / "backups"
    d.mkdir(parents=True, exist_ok=True)
    # 連続実行(二重送信など)で直前のバックアップを上書きしないよう、既存名は避ける
    stem = f"status_report_before_criteria_import_{datetime.now():%Y%m%d_%H%M%S_%f}"
    path = d / f"{stem}.db"
    n = 1
    while path.exists():
        path = d / f"{stem}_{n}.db"
        n += 1
    dst = sqlite3.connect(str(path))
    try:
        conn.backup(dst)
    finally:
        dst.close()
    return path


def apply(conn: sqlite3.Connection, preview: CriteriaPreview) -> dict:
    """プレビュー内容で基準を置換する。新規項目は無効状態で作成."""
    stats = {"updated": 0, "created": 0, "rows": 0}
    for plan in preview.plans:
        item_id = plan.item_id
        if item_id is None:
            spec = plan.new_item or {}
            after = repo.get_item_by_name(conn, spec.get("after") or "")
            item_id = repo.create_item(
                conn, name=plan.item_name, unit=spec.get("unit"), category=spec.get("category"),
                sort_order=after["sort_order"] if after else 0,
                direction=spec.get("direction", "higher_better"), in_radar=spec.get("in_radar", 0),
                source_column=spec.get("source_column"), derived_formula=None, is_qualitative=0,
            )
            repo.set_item_active(conn, item_id, False)
            stats["created"] += 1
        else:
            stats["updated"] += 1
        repo.delete_criteria_for_item(conn, item_id)
        for cs in plan.cond_sets:
            for thr, score, comment, desc in cs.rows:
                repo.add_criterion(conn, item_id, cs.sex, cs.age_band, cs.pacemaker, thr, score, comment, desc)
                stats["rows"] += 1
    return stats
