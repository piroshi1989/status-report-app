"""利用者 Excel 一括取込のテスト."""

from __future__ import annotations

import re
from datetime import datetime

import pytest
from openpyxl import Workbook

from app.services import patient_importer as pi

HEADERS = ["利用者コード", "氏名", "性別", "生年月日", "ペースメーカー", "服薬情報", "備考"]


@pytest.fixture()
def make_patients_xlsx(tmp_path):
    def _make(rows, headers=None):
        wb = Workbook()
        ws = wb.active
        ws.append(headers or HEADERS)
        for r in rows:
            ws.append(r)
        p = tmp_path / "patients.xlsx"
        wb.save(p)
        return p
    return _make


# --- パーサ単体 ------------------------------------------------------------

@pytest.mark.parametrize("raw,expected", [
    ("男", "M"), ("女", "F"), ("男性", "M"), ("女性", "F"),
    ("M", "M"), ("f", "F"), (1, "M"), (2, "F"), ("", None), ("不明", None),
])
def test_parse_sex(raw, expected):
    assert pi._parse_sex(raw) == expected


@pytest.mark.parametrize("raw,expected", [
    ("有", 1), ("あり", 1), ("○", 1), (1, 1), ("yes", 1),
    ("無", 0), ("なし", 0), ("×", 0), (0, 0),
    # 未入力は None(= 既存値を上書きしない)。0(無)と区別する。
    ("", None), (None, None),
])
def test_parse_pacemaker(raw, expected):
    assert pi._parse_pacemaker(raw) == expected


def test_parse_birth_variants():
    assert pi._parse_birth(datetime(1950, 5, 1))[0] == "1950-05-01"
    assert pi._parse_birth("1950-05-01")[0] == "1950-05-01"
    assert pi._parse_birth("1950/5/1")[0] == "1950-05-01"
    assert pi._parse_birth("1950年5月1日")[0] == "1950-05-01"
    # 未入力は正常扱い
    assert pi._parse_birth("") == (None, True)
    # 解釈不能
    assert pi._parse_birth("昭和25年")[1] is False


def test_parse_birth_excel_serial():
    # 1950-05-01 の Excel シリアル値
    serial = (datetime(1950, 5, 1) - pi._EXCEL_EPOCH).days
    assert pi._parse_birth(str(serial))[0] == "1950-05-01"


def test_parse_birth_handles_nat():
    import pandas as pd
    assert pi._parse_birth(pd.NaT) == (None, True)


# --- 取込フロー ------------------------------------------------------------

def test_import_new_patients(client, make_patients_xlsx):
    xlsx = make_patients_xlsx([
        ["U001", "山田太郎", "男", datetime(1950, 5, 1), "無", "", ""],
        ["U002", "佐藤花子", "女", "1948-11-23", "有", "血糖降下薬", "要見守り"],
    ])
    with open(xlsx, "rb") as f:
        r = client.post("/patients/import/analyze",
                        files={"file": ("p.xlsx", f, "application/octet-stream")})
    assert r.status_code == 200
    assert "山田太郎" in r.text
    token = re.search(r'name="token" value="([^"]+)"', r.text).group(1)

    r2 = client.post("/patients/import/commit",
                     data={"token": token, "update_existing": "1"},
                     follow_redirects=False)
    assert r2.status_code == 303
    assert "created=2" in r2.headers["location"]

    from app import db, repo
    with db.get_conn() as c:
        p = repo.get_patient_by_code(c, "U002")
        assert p["name"] == "佐藤花子"
        assert p["sex"] == "F"
        assert p["birth_date"] == "1948-11-23"
        assert p["pacemaker"] == 1
        assert p["med_info"] == "血糖降下薬"


def test_import_updates_existing_and_blank_does_not_overwrite(client, make_patients_xlsx):
    client.post("/patients/new", data={"code": "U010", "name": "旧名", "sex": "M",
                                       "birth_date": "1940-01-01", "pacemaker": "1",
                                       "med_info": "既存薬"})
    # 氏名だけ変更、他は空欄 → 空欄は既存値を保持
    xlsx = make_patients_xlsx([["U010", "新名", "", "", "", "", ""]])
    with open(xlsx, "rb") as f:
        r = client.post("/patients/import/analyze",
                        files={"file": ("p.xlsx", f, "application/octet-stream")})
    assert "更新" in r.text
    token = re.search(r'name="token" value="([^"]+)"', r.text).group(1)
    r2 = client.post("/patients/import/commit",
                     data={"token": token, "update_existing": "1"}, follow_redirects=False)
    assert "updated=1" in r2.headers["location"]

    from app import db, repo
    with db.get_conn() as c:
        p = repo.get_patient_by_code(c, "U010")
    assert p["name"] == "新名"
    assert p["sex"] == "M"             # 上書きされない
    assert p["birth_date"] == "1940-01-01"
    assert p["med_info"] == "既存薬"


def test_import_skips_existing_when_unchecked(client, make_patients_xlsx):
    client.post("/patients/new", data={"code": "U020", "name": "既存", "sex": "F"})
    xlsx = make_patients_xlsx([["U020", "上書き", "", "", "", "", ""],
                               ["U021", "新規", "男", "", "", "", ""]])
    with open(xlsx, "rb") as f:
        r = client.post("/patients/import/analyze",
                        files={"file": ("p.xlsx", f, "application/octet-stream")})
    token = re.search(r'name="token" value="([^"]+)"', r.text).group(1)
    r2 = client.post("/patients/import/commit", data={"token": token},  # update_existing なし
                     follow_redirects=False)
    loc = r2.headers["location"]
    assert "created=1" in loc and "updated=0" in loc and "skipped=1" in loc

    from app import db, repo
    with db.get_conn() as c:
        assert repo.get_patient_by_code(c, "U020")["name"] == "既存"


def test_import_flags_errors_and_duplicates(client, make_patients_xlsx):
    xlsx = make_patients_xlsx([
        ["E001", "不正性別", "ふめい", "", "", "", ""],
        ["E002", "不正日付", "男", "昭和25年", "", "", ""],
        ["E003", "重複1", "男", "", "", "", ""],
        ["E003", "重複2", "男", "", "", "", ""],
    ])
    with open(xlsx, "rb") as f:
        r = client.post("/patients/import/analyze",
                        files={"file": ("p.xlsx", f, "application/octet-stream")})
    assert "性別を解釈できません" in r.text
    assert "生年月日を解釈できません" in r.text
    assert "重複" in r.text

    token = re.search(r'name="token" value="([^"]+)"', r.text).group(1)
    r2 = client.post("/patients/import/commit", data={"token": token, "update_existing": "1"},
                     follow_redirects=False)
    loc = r2.headers["location"]
    # E003 の1件目は正常、2件目が重複エラー。E001/E002 はエラー。
    assert "created=1" in loc and "skipped=3" in loc


def test_import_requires_code_column(client, make_patients_xlsx):
    xlsx = make_patients_xlsx([["山田", "男"]], headers=["氏名", "性別"])
    with open(xlsx, "rb") as f:
        r = client.post("/patients/import/analyze",
                        files={"file": ("p.xlsx", f, "application/octet-stream")})
    assert r.status_code == 400
    assert "利用者コード" in r.text


def test_template_download(client):
    r = client.get("/patients/import/template")
    assert r.status_code == 200
    assert r.content[:2] == b"PK"  # xlsx = zip
    # テンプレートが自分自身で取り込めること
    import io
    from openpyxl import load_workbook
    ws = load_workbook(io.BytesIO(r.content)).active
    assert [c.value for c in ws[1]] == pi.TEMPLATE_HEADERS
