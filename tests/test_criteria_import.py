"""評価基準 Excel 取込(services/criteria_importer)と基準マスタ画面のテスト."""

from __future__ import annotations

import re
from pathlib import Path

import pytest

from app import config, db, repo
from app.services import criteria_importer as ci
from app.services.evaluation import evaluate
from app.services.report import build_report

HEADER = ["評価項目", "評価基準内容", "相対評価", "担当者", "評価コメント", "評価項目説明", "備考"]


def _grip_block(sex: str, top: float, comment_prefix: str) -> list[list]:
    low, mid = (8, 12) if top <= 20 else (15, 22)   # 90代は全体に低い(実ファイルと同じ傾向)
    return [
        [f"{sex}　握力(kg)", 0, 1, "OT", f"{comment_prefix}低い", "握力の説明", None],
        [None, low, 2, None, f"{comment_prefix}やや低い", None, None],
        [None, mid, 4, None, f"{comment_prefix}基準値", None, None],
        [None, top, 5, None, f"{comment_prefix}非常に優れる", None, None],
    ]


def _sheet_rows(band_label: str, grip_top: float, pm: bool) -> list[list]:
    rows = [[None, f"{band_label}基準値"], HEADER]
    rows += [
        ["HbA1c", 0, 5, None, "コントロールできています。", None, None],
        [None, 8.1, 2, None, "コントロールできていません。", None, None],
        [],
        ["血圧(mmHg)", "上", None, None, None, "上の血圧の説明", None],
        [None, 0, 5, None, "健康的な数値です。", None, None],
        [None, 180, 2, None, "要改善です。", None, None],
        [None, None, 1, None, "医師にご相談下さい。", None, None],   # 閾値が空欄
        [],
    ]
    rows += _grip_block("男", grip_top, f"{band_label}男")
    rows += [[]]
    rows += _grip_block("女", grip_top - 5, f"{band_label}女")
    rows += [[]]
    if pm:
        rows += [
            ["下腿周囲長(㎝)", "女", None, None, None, "下腿の説明", None],  # 性別が B 列の書式
            [None, 0, 1, None, "下腿 低い", None, None],
            [None, 33, 5, None, "下腿 良好", None, None],
            [],
        ]
    else:
        rows += [
            ["女　SMI", 0, 1, None, f"{band_label}SMI低い", None, None],
            [None, 5.8, 5, None, f"{band_label}SMI良好", None, None],
            [],
        ]
    rows += [
        ["FIM(社会性)", 1, "計測困難", None, "全面的に介助が必要です。", None, None],
        [None, 4, "修正自立", None, "適切な交流が出来ています。", None, None],
        [],
        ["(高音域)"],
        ["６０代"],
        ["聴力(4000Hz)", 0, 5, None, "よく聞こえています。", None, None],
        ["4000Hｚ", 20, 4, None, "同年代と同等です。", None, None],   # ブロック内の注記
        [None, 40, 3, None, "若干の衰えがあります。", None, None],
        [],
        ["体重減少率(%)", -100, 5, None, "正常範囲です。", None, None],
        [None, 1, 4, None, "気をつけましょう。", None, None],
        [None, 3, 2, None, "リスクが高くなっています。", None, None],
        [],
        ["視力", 1, 5, None, "非常に優れた視力です。", None, None],  # 降順に並んでいる
        [None, 0.3, 3, None, "平均値付近です。", None, None],
        [None, 0, 1, None, "眼科受診をお願いします。", None, None],
        [],
        ["歩行速度", 0, 1, None, "未対応の項目", None, None],
        [],
        ["総合評価", 1, "心身ともに良い状態です。", None, None, None, None],
        [None, 2, "改善に努めましょう。", None, None, None, None],
    ]
    return rows


@pytest.fixture()
def criteria_xlsx(tmp_path) -> Path:
    from openpyxl import Workbook

    wb = Workbook()
    wb.remove(wb.active)
    for title, label, top, pm in [
        ("50代以下", "50代", 46, False),
        ("70代 ", "70代", 41, False),
        ("70代 　ペースメーカー", "70代", 41, True),
        ("90代", "90代", 20, False),
        ("90代　ペースメーカー", "90代", 20, True),
    ]:
        ws = wb.create_sheet(title)
        for r in _sheet_rows(label, top, pm):
            ws.append(r)
    notes = wb.create_sheet("福水より申し送り")
    notes.append(["注意事項"])
    path = tmp_path / "評価項目 改訂案.xlsx"
    wb.save(path)
    return path


def _plan(preview, item_name):
    return next(p for p in preview.plans if p.item_name == item_name)


# --- シート名・正規化 ---------------------------------------------------------

def test_sheet_condition_parses_real_sheet_names():
    assert ci.sheet_condition("50代以下") == ("50", False)
    assert ci.sheet_condition("70代 ") == ("70", False)
    assert ci.sheet_condition("80代   ペースメーカー") == ("80", True)
    assert ci.sheet_condition("90代　ペースメーカー") == ("90", True)
    assert ci.sheet_condition("計測結果入力") is None


def test_normalize_does_not_spread_item_to_sheets_without_it():
    rows = [(0.0, 1, "a", None), (5.8, 5, "b", None)]
    all_conds = {("70", 0), ("70", 1), ("80", 0), ("80", 1)}
    # PM 無シートにしか無い項目は、全シート同一でも PM 有に広げない
    out = ci.normalize({("70", 0): rows, ("80", 0): rows}, all_conds)
    assert {(a, p) for a, p, _ in out} == {("70", 0), ("80", 0)}


def test_normalize_ignores_description_difference():
    a = [(0.0, 1, "c", "説明あり")]
    b = [(0.0, 1, "c", None)]
    assert ci.normalize({("60", 0): a, ("60", 1): b}) == [(None, None, a)]


# --- 解析(プレビュー) ----------------------------------------------------------

def test_analyze_maps_items_and_keeps_age_specific_comments(tmp_env, criteria_xlsx):
    db.init_db(seed=True)
    with db.get_conn() as conn:
        pv = ci.analyze(conn, criteria_xlsx)

    assert pv.skipped_sheets == ["福水より申し送り"]
    assert pv.unknown == ["歩行速度"]
    names = {p.item_name for p in pv.plans}
    assert {"HbA1c", "収縮期血圧", "握力", "SMI", "FIM", "聴力4000Hz",
            "体重増減率", "視力", "下腿周囲長"} <= names
    assert "総合評価" not in names

    # 全シート同一 → 条件なし 1 セット
    hba1c = _plan(pv, "HbA1c")
    assert [(c.sex, c.age_band, c.pacemaker) for c in hba1c.cond_sets] == [(None, None, None)]

    # 年代で違う握力は性別×年代ごとに保持し、PM 有無で同一なら PM 条件なし
    grip = _plan(pv, "握力")
    conds = {(c.sex, c.age_band, c.pacemaker) for c in grip.cond_sets}
    assert ("M", "70", None) in conds and ("M", "90", None) in conds
    by_cond = {(c.sex, c.age_band): c.rows for c in grip.cond_sets}
    assert by_cond[("M", "50")][-1][2] == "50代男非常に優れる"
    assert by_cond[("M", "90")][-1][:3] == (20.0, 5, "90代男非常に優れる")


def test_analyze_handles_excel_layout_quirks(tmp_env, criteria_xlsx):
    db.init_db(seed=True)
    with db.get_conn() as conn:
        pv = ci.analyze(conn, criteria_xlsx)

    # 血圧: 「上」の小見出しは読み飛ばし、閾値空欄の行は警告して除外
    bp = _plan(pv, "収縮期血圧")
    assert [(r[0], r[1]) for r in bp.cond_sets[0].rows] == [(0.0, 5), (180.0, 2)]
    assert any("血圧(mmHg)" in w and "空欄" in w for w in pv.warnings)

    # 「4000Hｚ」の注記行もブロックの一部として読む
    hearing = _plan(pv, "聴力4000Hz")
    assert [r[0] for r in hearing.cond_sets[0].rows] == [0.0, 20.0, 40.0]

    # 定性項目はコード値が評価、コメントは評価コメント列
    fim = _plan(pv, "FIM")
    assert [(r[0], r[1], r[2]) for r in fim.cond_sets[0].rows] == [
        (1.0, 1, "全面的に介助が必要です。"), (4.0, 4, "適切な交流が出来ています。"),
    ]

    # 降順に書かれた視力は閾値昇順に並べ替える(視力は seed 済みの無効項目)
    eye = _plan(pv, "視力")
    assert not eye.is_new
    assert [r[0] for r in eye.cond_sets[0].rows] == [0.0, 0.3, 1.0]

    # 性別が B 列にある下腿周囲長は女性・PM 有の基準になる(PM 無シートには無い)
    calf = _plan(pv, "下腿周囲長")
    assert {(c.sex, c.pacemaker) for c in calf.cond_sets} == {("F", 1)}

    # SMI は PM 無シートにしか無いので PM 有の患者へ広げない
    smi = _plan(pv, "SMI")
    assert {c.pacemaker for c in smi.cond_sets} <= {0, None}
    assert all(c.age_band is not None for c in smi.cond_sets)


# --- 適用 ---------------------------------------------------------------------

def test_apply_replaces_criteria_and_creates_inactive_items(tmp_env, criteria_xlsx):
    db.init_db(seed=True)
    with db.get_conn() as conn:
        # seed に無い古い DB を想定して、Excel 由来の項目を消しておく(取込で無効のまま作られる)
        conn.execute("DELETE FROM eval_items WHERE name IN ('視力', '下腿周囲長')")
        one_leg = repo.get_item_by_name(conn, "片脚立位時間")
        before_one_leg = len(repo.list_criteria_for_item(conn, one_leg["id"]))
        pv = ci.analyze(conn, criteria_xlsx)
        stats = ci.apply(conn, pv)

        assert stats["created"] == 2
        eye = repo.get_item_by_name(conn, "視力")
        assert eye["is_active"] == 0
        assert eye["sort_order"] == repo.get_item_by_name(conn, "聴力4000Hz")["sort_order"]
        # Excel に無い項目の基準は変えない
        assert len(repo.list_criteria_for_item(conn, one_leg["id"])) == before_one_leg

        grip_id = repo.get_item_by_name(conn, "握力")["id"]
        rows = repo.criterion_rows_for_item(conn, grip_id)
        # 同じ握力 25kg・男性でも年代でコメント(と評価)が変わる
        r70 = evaluate(rows, 25, sex="M", age_band="70", pacemaker=0)
        r90 = evaluate(rows, 25, sex="M", age_band="90", pacemaker=1)
        assert (r70.score, r70.comment) == (4, "70代男基準値")
        assert (r90.score, r90.comment) == (5, "90代男非常に優れる")


def test_import_screen_flow_with_backup(client, criteria_xlsx):
    assert client.get("/criteria/import").status_code == 200
    with open(criteria_xlsx, "rb") as f:
        r = client.post("/criteria/import/analyze", files={"file": ("基準.xlsx", f, "application/octet-stream")})
    assert r.status_code == 200
    assert "歩行速度" in r.text           # 未対応項目の表示
    assert "視力" in r.text                  # seed 済みの無効項目にも取り込む
    assert "90代男非常に優れる" in r.text   # 条件ごとの詳細
    token = re.search(r'name="token" value="([^"]+)"', r.text).group(1)

    r2 = client.post("/criteria/import/commit", data={"token": token}, follow_redirects=False)
    assert r2.status_code == 303
    assert r2.headers["location"].startswith("/criteria?imported=")
    backups = list((config.data_dir() / "backups").glob("*.db"))
    assert len(backups) == 1

    lst = client.get(r2.headers["location"])
    assert "評価基準を取り込みました" in lst.text
    assert backups[0].name in lst.text


def test_import_rejects_xls(client, tmp_path):
    bogus = tmp_path / "old.xls"
    bogus.write_bytes(b"not really xls")
    with open(bogus, "rb") as f:
        r = client.post("/criteria/import/analyze", files={"file": ("old.xls", f, "application/vnd.ms-excel")})
    assert r.status_code == 400
    assert ".xlsx" in r.text


def test_report_comment_changes_with_age_band(client, criteria_xlsx):
    """取込後、同じ測定値でも年代で報告書のコメントが変わる."""
    with db.get_conn() as conn:
        ci.apply(conn, ci.analyze(conn, criteria_xlsx))
        grip_id = repo.get_item_by_name(conn, "握力")["id"]
        sid = repo.create_session(conn, 1, "2026-07-01", "m.xlsx")
        lines = {}
        for code, birth in [("A", "1953-01-01"), ("B", "1933-01-01")]:   # 73歳 / 93歳
            pid = repo.create_patient(conn, code=code, name=code, sex="M", birth_date=birth)
            repo.upsert_measurement(conn, sid, pid, grip_id, 25.0, "25")
            data = build_report(conn, pid, sid)
            lines[code] = next(ln for ln in data.lines if ln.name == "握力")
    assert lines["A"].comment == "70代男基準値"
    assert lines["B"].comment == "90代男非常に優れる"


def test_weight_change_scored_by_loss_rate(client, criteria_xlsx):
    with db.get_conn() as conn:
        ci.apply(conn, ci.analyze(conn, criteria_xlsx))
        wid = repo.get_item_by_name(conn, "体重")["id"]
        pid = repo.create_patient(conn, code="W", name="W", sex="F", birth_date="1950-01-01")
        s1 = repo.create_session(conn, 1, "2026-01-01", "a.xlsx")
        s2 = repo.create_session(conn, 2, "2026-02-01", "b.xlsx")
        repo.upsert_measurement(conn, s1, pid, wid, 50.0, "50")
        repo.upsert_measurement(conn, s2, pid, wid, 49.5, "49.5")   # 減少率 1% ちょうど
        line = next(ln for ln in build_report(conn, pid, s2).lines if ln.name == "体重増減率")
    assert line.raw_text == "99%(1%減少)"
    assert (line.score, line.comment) == (4, "気をつけましょう。")


# --- 有効チェック ---------------------------------------------------------------

def _item_form(item) -> dict:
    return {
        "name": item["name"], "unit": item["unit"] or "", "category": item["category"] or "",
        "sort_order": str(item["sort_order"]), "direction": item["direction"],
        "source_column": item["source_column"] or "", "derived_formula": item["derived_formula"] or "",
        "in_radar": "1",
    }


def test_unchecking_active_disables_item(client):
    with db.get_conn() as conn:
        item = repo.get_item_by_name(conn, "片脚立位時間")
    # チェックを外すとブラウザは is_active を送らない
    r = client.post(f"/criteria/{item['id']}/item", data=_item_form(item), follow_redirects=False)
    assert r.status_code == 303
    with db.get_conn() as conn:
        assert repo.get_item(conn, item["id"])["is_active"] == 0
        assert "片脚立位時間" not in {it["name"] for it in repo.list_items(conn)}

    r = client.post(f"/criteria/{item['id']}/item", data={**_item_form(item), "is_active": "1"},
                    follow_redirects=False)
    with db.get_conn() as conn:
        assert repo.get_item(conn, item["id"])["is_active"] == 1


# --- 実ファイル(リポジトリ直下に置かれている場合のみ) ------------------------------

REAL_FILES = sorted(Path(__file__).resolve().parent.parent.glob("評価項目*.xlsx"))


@pytest.mark.skipif(not REAL_FILES, reason="評価基準の実 Excel が無い")
@pytest.mark.parametrize("path", REAL_FILES, ids=lambda p: p.stem)
def test_real_workbook_maps_every_item(tmp_env, path):
    db.init_db(seed=True)
    with db.get_conn() as conn:
        pv = ci.analyze(conn, path)
    assert pv.unknown == []
    # 最終版は CRP が増えている
    assert len(pv.plans) == (21 if "最終版" in path.name else 20)
    assert all(not p.is_new for p in pv.plans)   # Excel の項目はすべて seed 済み
    grip = _plan(pv, "握力")
    tops = {(c.sex, c.age_band): c.rows[-1][0] for c in grip.cond_sets}
    assert tops[("M", "50")] == 46 and tops[("M", "80")] == 32


# --- レビュー指摘の回帰テスト -----------------------------------------------------

def _tug_workbook(tmp_path, spellings: dict[str, str]) -> Path:
    from openpyxl import Workbook

    wb = Workbook()
    wb.remove(wb.active)
    for title, name in spellings.items():
        ws = wb.create_sheet(title)
        for r in [[None, f"{title}基準値"], HEADER,
                  [name, 0, 1, None, f"{title}遅い", None, None],
                  [None, 12, 4, None, f"{title}基準値", None, None]]:
            ws.append(r)
    path = tmp_path / "tug.xlsx"
    wb.save(path)
    return path


def test_spelling_variants_merge_into_one_item(tmp_env, tmp_path):
    """シートにより「TUG(秒)」「TUG」と表記がゆれても、どの年代の基準も消えない."""
    path = _tug_workbook(tmp_path, {"60代": "TUG(秒)", "70代": "TUG", "80代": "TUG(秒)"})
    db.init_db(seed=True)
    with db.get_conn() as conn:
        pv = ci.analyze(conn, path)
        assert [p.item_name for p in pv.plans] == ["TUG"]
        assert pv.plans[0].excel_name == "TUG(秒) / TUG"
        stats = ci.apply(conn, pv)
        assert stats["updated"] == 1
        bands = {r["age_band"] for r in repo.list_criteria_for_item(conn, pv.plans[0].item_id)}
    assert bands == {"60", "70", "80"}


def test_backup_never_overwrites_previous(tmp_env):
    db.init_db(seed=True)
    with db.get_conn() as conn:
        paths = {ci.backup_db(conn) for _ in range(5)}
    assert len(paths) == 5 and all(p.exists() for p in paths)


def test_unknown_name_inside_block_is_warned():
    rows = [HEADER,
            ["TUG(秒)", 0, 1, None, "a", None, None],
            [None, 12, 4, None, "b", None, None],
            ["歩行速度", 0, 1, None, "x", None, None],      # 空行なしで続いた未知の項目
            [None, 12.5, 4.5, None, "y", None, None]]       # 小数の評価
    parsed, warnings = ci.parse_sheet(rows, "60代")
    assert parsed[("TUG(秒)", None)] == [(0.0, 1, "a", None), (12.0, 4, "b", None)]
    assert any("歩行速度" in w and "一部として" in w for w in warnings)
    assert any("重複" in w for w in warnings)
    assert any("整数でない" in w for w in warnings)


def test_migrate_dry_run_does_not_create_db(tmp_env, criteria_xlsx, capsys):
    from scripts.migrate_excel import migrate

    migrate(criteria_xlsx, dry_run=True)
    assert "握力" in capsys.readouterr().out
    assert not config.db_path().exists()
